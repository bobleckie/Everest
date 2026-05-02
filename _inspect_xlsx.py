"""Dump the pricing workbooks so we can design the page."""
import openpyxl
import sys
sys.stdout.reconfigure(encoding="utf-8")


def dump_sheet(ws, max_rows=None, max_cols=None):
    max_rows = max_rows or ws.max_row
    max_cols = max_cols or ws.max_column
    for row in ws.iter_rows(min_row=1, max_row=max_rows, max_col=max_cols, values_only=False):
        parts = []
        for c in row:
            v = c.value
            if v is None or (isinstance(v, str) and not v.strip()):
                continue
            if isinstance(v, float):
                v = round(v, 4)
            parts.append(f"{c.coordinate}={v!r}")
        if parts:
            print("  " + " | ".join(parts))


def dump_file(path, max_cells=400):
    print("=" * 100)
    print("FILE:", path)
    print("=" * 100)
    wb = openpyxl.load_workbook(path, data_only=True)
    for s in wb.sheetnames:
        ws = wb[s]
        print(f"\n---- SHEET: {s}  ({ws.max_row} rows x {ws.max_column} cols) ----")
        dump_sheet(ws)


dump_file(r"docs\Copy of NJ 2019 v1.20.xlsx")
print("\n\n")
dump_file(r"docs\NJ Vendor Cost Comparison 12-12-23.xlsx")
