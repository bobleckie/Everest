from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from jinja2 import Template
from docx import Document
from pptx import Presentation
import io

def generate_pdf(content: str) -> bytes:
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=letter)
    c.drawString(100, 750, content[:1000])  # Simple text
    c.save()
    buffer.seek(0)
    return buffer.getvalue()

def generate_html(content: str) -> str:
    template = Template("<html><body><h1>RFP Response</h1><p>{{ content }}</p></body></html>")
    return template.render(content=content)

def generate_word(content: str) -> bytes:
    doc = Document()
    doc.add_heading('RFP Response', 0)
    doc.add_paragraph(content)
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def generate_ppt(content: str) -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    title = slide.shapes.title
    title.text = "RFP Response"
    body = slide.placeholders[1]
    body.text = content[:500]
    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()