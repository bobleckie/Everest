"""
Schedule / Key Events router.

Endpoints:
  GET    /api/schedule/events                  — list events (with filters)
  POST   /api/schedule/events                  — manually create an event
  PUT    /api/schedule/events/{id}             — update an event
  DELETE /api/schedule/events/{id}             — delete an event
  POST   /api/schedule/extract/{doc_id}        — AI-extract schedule from a doc
  POST   /api/schedule/derive/{proposal_id}    — rebuild internal deadlines
  GET    /api/schedule/upcoming                — next N upcoming events
"""
from __future__ import annotations

import logging
from datetime import datetime
from typing import List, Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from ..auth import get_current_user
from ..database import get_db
from ..models import IngestedDocument, Proposal, RfpScheduleEvent, User
from ..services.schedule_extractor import (
    derive_internal_deadlines,
    extract_schedule_events,
    _resolve_relative_dates,
)

logger = logging.getLogger(__name__)
router = APIRouter()


# ── Schemas ──────────────────────────────────────────────────────────

class ScheduleEventOut(BaseModel):
    id: int
    document_id: Optional[int]
    proposal_id: Optional[int]
    event_type: str
    label: str
    event_date: Optional[datetime]
    end_date: Optional[datetime]
    is_mandatory: bool
    source_page: Optional[int]
    source_text: Optional[str]
    notes: Optional[str]
    confidence: str
    extracted_by_ai: bool
    verified: bool
    verified_by: Optional[str]
    assignee: Optional[str]
    status: str
    created_at: Optional[datetime]
    updated_at: Optional[datetime]

    class Config:
        from_attributes = True


class ScheduleEventCreate(BaseModel):
    event_type: str
    label: str
    event_date: Optional[datetime] = None
    end_date: Optional[datetime] = None
    is_mandatory: bool = False
    document_id: Optional[int] = None
    proposal_id: Optional[int] = None
    source_page: Optional[int] = None
    source_text: Optional[str] = None
    notes: Optional[str] = None
    confidence: str = "high"
    assignee: Optional[str] = None
    status: str = "upcoming"


class ScheduleEventUpdate(BaseModel):
    event_type: Optional[str] = None
    label: Optional[str] = None
    event_date: Optional[datetime] = None
    end_date: Optional[datetime] = None
    is_mandatory: Optional[bool] = None
    source_page: Optional[int] = None
    source_text: Optional[str] = None
    notes: Optional[str] = None
    confidence: Optional[str] = None
    verified: Optional[bool] = None
    assignee: Optional[str] = None
    status: Optional[str] = None


# ── Helpers ──────────────────────────────────────────────────────────

def _to_dict(ev: RfpScheduleEvent, doc_name: Optional[str] = None) -> dict:
    return {
        "id": ev.id,
        "document_id": ev.document_id,
        "document_name": doc_name,
        "proposal_id": ev.proposal_id,
        "event_type": ev.event_type,
        "label": ev.label,
        "event_date": ev.event_date.isoformat() if ev.event_date else None,
        "end_date": ev.end_date.isoformat() if ev.end_date else None,
        "is_mandatory": bool(ev.is_mandatory),
        "source_page": ev.source_page,
        "source_text": ev.source_text,
        "notes": ev.notes,
        "confidence": ev.confidence,
        "extracted_by_ai": bool(ev.extracted_by_ai),
        "verified": bool(ev.verified),
        "verified_by": ev.verified_by,
        "assignee": ev.assignee,
        "status": ev.status,
        # Relative-deadline fields (#401, 2026-04-28).
        "offset_days": getattr(ev, "offset_days", None),
        "offset_anchor": getattr(ev, "offset_anchor", None),
        "is_draft_with_quote": bool(getattr(ev, "is_draft_with_quote", False)),
        "section_refs": getattr(ev, "section_refs", None),
        "event_date_resolved": bool(getattr(ev, "event_date_resolved", False)),
        "created_at": ev.created_at.isoformat() if ev.created_at else None,
        "updated_at": ev.updated_at.isoformat() if ev.updated_at else None,
    }


# ── List / Get ───────────────────────────────────────────────────────

@router.get("/events")
def list_events(
    document_id: Optional[int] = Query(None),
    proposal_id: Optional[int] = Query(None),
    event_type: Optional[str] = Query(None),
    extracted_by_ai: Optional[bool] = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    q = db.query(RfpScheduleEvent)
    if document_id is not None:
        q = q.filter(RfpScheduleEvent.document_id == document_id)
    if proposal_id is not None:
        q = q.filter(RfpScheduleEvent.proposal_id == proposal_id)
    if event_type:
        q = q.filter(RfpScheduleEvent.event_type == event_type)
    if extracted_by_ai is not None:
        q = q.filter(RfpScheduleEvent.extracted_by_ai == extracted_by_ai)

    # Order: events with dates first (by date asc), then undated events
    rows = q.all()
    rows.sort(key=lambda r: (r.event_date is None, r.event_date or datetime.max, r.id))

    # Batch-load document names so each event can show its source doc
    doc_ids = {r.document_id for r in rows if r.document_id is not None}
    doc_names: dict = {}
    if doc_ids:
        for d in db.query(IngestedDocument).filter(IngestedDocument.id.in_(doc_ids)).all():
            doc_names[d.id] = d.original_filename or d.filename

    return {
        "events": [_to_dict(r, doc_name=doc_names.get(r.document_id)) for r in rows],
        "count": len(rows),
    }


@router.get("/upcoming")
def upcoming_events(
    limit: int = Query(5, ge=1, le=50),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    now = datetime.utcnow()
    rows = (
        db.query(RfpScheduleEvent)
        .filter(RfpScheduleEvent.event_date.isnot(None))
        .filter(RfpScheduleEvent.event_date >= now)
        .filter(RfpScheduleEvent.status != "cancelled")
        .order_by(RfpScheduleEvent.event_date.asc())
        .limit(limit)
        .all()
    )
    doc_ids = {r.document_id for r in rows if r.document_id is not None}
    doc_names: dict = {}
    if doc_ids:
        for d in db.query(IngestedDocument).filter(IngestedDocument.id.in_(doc_ids)).all():
            doc_names[d.id] = d.original_filename or d.filename
    return {"events": [_to_dict(r, doc_name=doc_names.get(r.document_id)) for r in rows]}


# ── Create / Update / Delete ─────────────────────────────────────────

@router.post("/events")
def create_event(
    body: ScheduleEventCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if body.document_id is not None:
        if not db.query(IngestedDocument).filter(IngestedDocument.id == body.document_id).first():
            raise HTTPException(404, f"Document {body.document_id} not found")
    if body.proposal_id is not None:
        if not db.query(Proposal).filter(Proposal.id == body.proposal_id).first():
            raise HTTPException(404, f"Proposal {body.proposal_id} not found")

    ev = RfpScheduleEvent(
        document_id=body.document_id,
        proposal_id=body.proposal_id,
        event_type=body.event_type,
        label=body.label[:250],
        event_date=body.event_date,
        end_date=body.end_date,
        is_mandatory=body.is_mandatory,
        source_page=body.source_page,
        source_text=(body.source_text or "")[:1000],
        notes=(body.notes or "")[:2000],
        confidence=body.confidence,
        extracted_by_ai=False,
        verified=True,
        verified_by=current_user.username,
        assignee=body.assignee,
        status=body.status,
    )
    db.add(ev)
    db.commit()
    db.refresh(ev)
    doc_name = None
    if ev.document_id:
        d = db.query(IngestedDocument).filter(IngestedDocument.id == ev.document_id).first()
        if d:
            doc_name = d.original_filename or d.filename
    return _to_dict(ev, doc_name=doc_name)


@router.put("/events/{event_id}")
def update_event(
    event_id: int,
    body: ScheduleEventUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    ev = db.query(RfpScheduleEvent).filter(RfpScheduleEvent.id == event_id).first()
    if not ev:
        raise HTTPException(404, f"Event {event_id} not found")

    payload = body.model_dump(exclude_unset=True)
    for k, v in payload.items():
        setattr(ev, k, v)

    if body.verified and not ev.verified_by:
        ev.verified_by = current_user.username

    db.commit()
    db.refresh(ev)
    doc_name = None
    if ev.document_id:
        d = db.query(IngestedDocument).filter(IngestedDocument.id == ev.document_id).first()
        if d:
            doc_name = d.original_filename or d.filename
    return _to_dict(ev, doc_name=doc_name)


@router.delete("/events/{event_id}")
def delete_event(
    event_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    ev = db.query(RfpScheduleEvent).filter(RfpScheduleEvent.id == event_id).first()
    if not ev:
        raise HTTPException(404, f"Event {event_id} not found")
    db.delete(ev)
    db.commit()
    return {"deleted": event_id}


# ── Extraction ───────────────────────────────────────────────────────

@router.post("/extract/{document_id}")
def extract_from_document(
    document_id: int,
    proposal_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    doc = db.query(IngestedDocument).filter(IngestedDocument.id == document_id).first()
    if not doc:
        raise HTTPException(404, f"Document {document_id} not found")

    try:
        result = extract_schedule_events(
            db,
            document_id=document_id,
            proposal_id=proposal_id,
            replace_existing=True,
        )
    except Exception as e:
        logger.exception(f"Schedule extraction failed for doc {document_id}")
        raise HTTPException(500, f"Extraction failed: {e}")

    return result


@router.post("/derive/{proposal_id}")
def derive_for_proposal(
    proposal_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Regenerate internal working-backwards deadlines for a proposal."""
    p = db.query(Proposal).filter(Proposal.id == proposal_id).first()
    if not p:
        raise HTTPException(404, f"Proposal {proposal_id} not found")
    if not p.due_date:
        raise HTTPException(400, "Proposal has no due_date set — extract schedule first or set it manually")
    count = derive_internal_deadlines(db, proposal_id)
    return {"proposal_id": proposal_id, "derived": count}


# ── Contract anchor (the "Start" date for relative deadlines) ────────

class ContractAnchorBody(BaseModel):
    contract_effective_date: Optional[datetime] = None


@router.get("/proposal/{proposal_id}/contract-anchor")
def get_contract_anchor(
    proposal_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    p = db.query(Proposal).filter(Proposal.id == proposal_id).first()
    if not p:
        raise HTTPException(404, f"Proposal {proposal_id} not found")
    ced = getattr(p, "contract_effective_date", None)
    relative_count = (
        db.query(RfpScheduleEvent)
        .filter(
            RfpScheduleEvent.proposal_id == proposal_id,
            RfpScheduleEvent.offset_days.isnot(None),
        )
        .count()
    )
    return {
        "proposal_id": proposal_id,
        "contract_effective_date": ced.isoformat() if ced else None,
        "relative_event_count": relative_count,
    }


@router.put("/proposal/{proposal_id}/contract-anchor")
def set_contract_anchor(
    proposal_id: int,
    body: ContractAnchorBody,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Set or clear `proposals.contract_effective_date`. When set, every
    relative deadline (offset_days != NULL) gets its absolute event_date
    recomputed. Clearing reverts only previously-resolved rows back to
    NULL event_date (user-entered dates are never touched)."""
    p = db.query(Proposal).filter(Proposal.id == proposal_id).first()
    if not p:
        raise HTTPException(404, f"Proposal {proposal_id} not found")
    p.contract_effective_date = body.contract_effective_date
    db.commit()

    if body.contract_effective_date is None:
        # Reverting: blank out only resolver-managed event_dates.
        db.query(RfpScheduleEvent).filter(
            RfpScheduleEvent.proposal_id == proposal_id,
            RfpScheduleEvent.event_date_resolved == True,  # noqa: E712
            RfpScheduleEvent.offset_anchor == "contract_start",
        ).update(
            {"event_date": None, "event_date_resolved": False},
            synchronize_session=False,
        )
        db.commit()
        return {
            "proposal_id": proposal_id,
            "contract_effective_date": None,
            "resolved": 0,
            "cleared": True,
        }

    n = _resolve_relative_dates(db, proposal_id)
    return {
        "proposal_id": proposal_id,
        "contract_effective_date": p.contract_effective_date.isoformat(),
        "resolved": n,
        "cleared": False,
    }


# ── Deduplicate ──────────────────────────────────────────────────────

@router.post("/dedupe")
def dedupe_events(
    proposal_id: Optional[int] = Query(None, description="Restrict to one proposal"),
    apply: bool = Query(False, description="If false (default), only report what would be removed"),
    same_day_window_hours: int = Query(48, ge=0, le=168,
        description="Two events of the same type within this many hours collapse to one"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Collapse near-duplicate schedule events.

    Two events are considered duplicates of each other when ALL of:
      * same proposal_id (or both NULL)
      * same event_type
      * either (a) both have event_date and are within
        ``same_day_window_hours`` of each other, OR (b) both have NULL
        event_date and the same label
      * not in a terminal state where the user already declared one
        canonical (we keep complete/cancelled rows untouched)

    Among a duplicate group we keep the row that scores highest on:
      verified > extracted_by_ai=False > confidence (high>medium>low) >
      longer source_text > most recent updated_at > lowest id

    Defaults to DRY-RUN — pass ``?apply=true`` to actually delete.
    """
    from datetime import timedelta as _td
    q = db.query(RfpScheduleEvent)
    if proposal_id is not None:
        q = q.filter(RfpScheduleEvent.proposal_id == proposal_id)
    rows = q.all()

    # Score function — higher is better, we keep the highest in each group.
    _conf_score = {"high": 3, "medium": 2, "low": 1}

    def score(r: RfpScheduleEvent) -> tuple:
        return (
            1 if r.verified else 0,
            0 if r.extracted_by_ai else 1,  # manual / derived beats AI
            _conf_score.get((r.confidence or "").lower(), 0),
            len(r.source_text or ""),
            r.updated_at.timestamp() if r.updated_at else 0,
            -r.id,
        )

    # Group: (proposal_id, event_type) -> rows. Then bucket within group
    # by date proximity.
    from collections import defaultdict
    by_key: dict = defaultdict(list)
    for r in rows:
        # Skip terminal-state events — user has already curated those.
        if (r.status or "").lower() in ("complete", "cancelled"):
            continue
        by_key[(r.proposal_id, r.event_type)].append(r)

    to_delete: list[int] = []
    groups_collapsed = 0
    window = _td(hours=same_day_window_hours)

    for key, items in by_key.items():
        if len(items) < 2:
            continue
        # Sort by date so the sliding-window grouping is deterministic.
        items_sorted = sorted(items, key=lambda r: (
            r.event_date or datetime.max, r.id,
        ))
        # Build groups where consecutive items either share NULL date+label
        # or have event_dates within `window`.
        bucket: list[RfpScheduleEvent] = []
        buckets: list[list[RfpScheduleEvent]] = []
        for r in items_sorted:
            if not bucket:
                bucket = [r]
                continue
            prev = bucket[-1]
            same_bucket = False
            if prev.event_date and r.event_date:
                same_bucket = abs(r.event_date - prev.event_date) <= window
            elif (prev.event_date is None) and (r.event_date is None):
                same_bucket = (prev.label or "").strip().lower() == (r.label or "").strip().lower()
            if same_bucket:
                bucket.append(r)
            else:
                buckets.append(bucket)
                bucket = [r]
        if bucket:
            buckets.append(bucket)

        for b in buckets:
            if len(b) < 2:
                continue
            groups_collapsed += 1
            # Pick the keeper with the best score; mark the rest for deletion.
            keeper = max(b, key=score)
            for r in b:
                if r.id != keeper.id:
                    to_delete.append(r.id)

    result = {
        "dry_run": not apply,
        "considered": len(rows),
        "groups_collapsed": groups_collapsed,
        "would_remove": len(to_delete),
        "ids": to_delete[:100],  # cap for response size
    }
    if apply and to_delete:
        (
            db.query(RfpScheduleEvent)
            .filter(RfpScheduleEvent.id.in_(to_delete))
            .delete(synchronize_session=False)
        )
        db.commit()
        result["removed"] = len(to_delete)
    return result


# ── Export ───────────────────────────────────────────────────────────

# Human-readable type labels, kept in sync with the React TYPE_LABEL map.
_TYPE_LABEL = {
    "question_period_start": "Q&A Period Opens",
    "question_period_end": "Questions Due",
    "pre_bid_conference": "Pre-Bid Conference",
    "mandatory_site_visit": "Site Visit",
    "addenda_cutoff": "Addenda Cutoff",
    "proposal_due": "Proposal Due",
    "bid_opening": "Bid Opening",
    "evaluation_period_start": "Evaluation Begins",
    "evaluation_period_end": "Evaluation Ends",
    "bafo_due": "BAFO Due",
    "award_notification": "Award Notification",
    "contract_start": "Contract Start",
    "period_of_performance_start": "PoP Start",
    "period_of_performance_end": "PoP End",
    "implementation_milestone": "Implementation Milestone",
    "internal_review": "Internal Review",
    "internal_sme_signoff": "SME Sign-Off",
    "internal_red_team": "Red-Team Review",
    "internal_pricing_final": "Pricing Final",
    "internal_production": "Production & Packaging",
    "other": "Other",
}


@router.get("/export.xlsx")
def export_schedule_xlsx(
    document_id: Optional[int] = Query(None),
    proposal_id: Optional[int] = Query(None),
    event_type: Optional[str] = Query(None),
    extracted_by_ai: Optional[bool] = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Export the filtered schedule to an XLSX file.

    Accepts the same filters as `GET /events`. Rows are ordered chronologically,
    with undated events at the bottom.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except Exception as e:
        raise HTTPException(500, f"openpyxl not installed: {e}")
    from fastapi.responses import StreamingResponse
    from io import BytesIO

    q = db.query(RfpScheduleEvent)
    if document_id is not None:
        q = q.filter(RfpScheduleEvent.document_id == document_id)
    if proposal_id is not None:
        q = q.filter(RfpScheduleEvent.proposal_id == proposal_id)
    if event_type:
        q = q.filter(RfpScheduleEvent.event_type == event_type)
    if extracted_by_ai is not None:
        q = q.filter(RfpScheduleEvent.extracted_by_ai == extracted_by_ai)

    rows = q.all()
    rows.sort(key=lambda r: (r.event_date is None, r.event_date or datetime.max, r.id))

    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"

    headers = [
        "ID",
        "Event Type",
        "Label",
        "Event Date",
        "End Date",
        "Mandatory",
        "Status",
        "Confidence",
        "Assignee",
        "Source Doc ID",
        "Source Page",
        "AI-Extracted",
        "Verified",
        "Verified By",
        "Notes",
        "Source Text",
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="00AEE6", end_color="00AEE6", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def _fmt_dt(d):
        if not d:
            return ""
        try:
            # Naïve datetime, local/UTC-agnostic — display as ISO minute precision.
            return d.strftime("%Y-%m-%d %H:%M")
        except Exception:
            return str(d)

    for ev in rows:
        ws.append([
            ev.id,
            _TYPE_LABEL.get(ev.event_type, ev.event_type or ""),
            ev.label or "",
            _fmt_dt(ev.event_date),
            _fmt_dt(ev.end_date),
            "Yes" if ev.is_mandatory else "No",
            ev.status or "",
            ev.confidence or "",
            ev.assignee or "",
            ev.document_id if ev.document_id is not None else "",
            ev.source_page if ev.source_page is not None else "",
            "Yes" if ev.extracted_by_ai else "No",
            "Yes" if ev.verified else "No",
            ev.verified_by or "",
            (ev.notes or "")[:32000],
            (ev.source_text or "")[:32000],
        ])

    widths = [6, 22, 40, 18, 18, 10, 12, 11, 16, 13, 11, 13, 10, 14, 50, 60]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"
    # Wrap-text on the long free-form columns.
    for row_idx in range(2, ws.max_row + 1):
        for col_idx in (3, 15, 16):
            ws.cell(row=row_idx, column=col_idx).alignment = Alignment(
                vertical="top", wrap_text=True
            )

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    stamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    filename = f"schedule_{stamp}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
